"""Generate the sample team list samples/team/team_list.xlsx (fictional data).

Needs openpyxl (requirements-dev.txt). Usage: python samples/make_team.py
"""
import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

TODAY = dt.date(2026, 9, 29)
IN_TWO_WEEKS = TODAY + dt.timedelta(days=14)
YEAR_END = dt.date(2026, 12, 31)
NEXT_YEAR = dt.date(2027, 1, 1)

# (name as in HR list, level, location, current assignment, assignment end, availability, available from)
ROWS = [
    ("Markus Mustermann", "Manager", "Zurich", "Swiss Utilities Company - Architecture", YEAR_END, 0.0, NEXT_YEAR),
    ("Manuela Mustermann", "Associate Manager", "Zurich", "Smart Meter Platform - Team Lead", YEAR_END, 0.0, NEXT_YEAR),
    ("Sophie Marie Keller", "Senior Consultant", "Zurich", "Insurance Claims Portal", YEAR_END, 0.0, NEXT_YEAR),
    ("Lukas Brändle", "Consultant", "Bern", "Telecom Data Platform", YEAR_END, 0.0, NEXT_YEAR),
    ("Aylin Demir", "Consultant", "Basel", "", None, 1.0, TODAY),
    ("Tomasz Nowak", "Consultant", "Zurich", "Payments Modernisation", IN_TWO_WEEKS - dt.timedelta(days=1), 0.0, IN_TWO_WEEKS),
    ("Elena Rossi", "Senior Consultant", "Lucerne", "", None, 1.0, TODAY),
    ("David Okafor", "Senior Consultant", "Basel", "Pharma LIMS Integration", YEAR_END, 0.0, NEXT_YEAR),
    ("Julia Hofmann", "Analyst", "Zurich", "Broker Portal", YEAR_END, 0.0, NEXT_YEAR),
    ("Nikolai Petrov", "Senior Consultant", "Zurich", "", None, 1.0, TODAY),
]

wb = Workbook()
ws = wb.active
ws.title = "Team"
ws.append(["Team Engineering Zurich - staffing overview (sample data)"])
ws["A1"].font = Font(bold=True, size=13)
ws.append([f"As of {TODAY:%d.%m.%Y}"])
ws.append([])
headers = ["Name", "Career Level", "Location", "Current Assignment", "Assignment End", "Availability", "Available From"]
ws.append(headers)
for c in ws[4]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="2F3C7E")
    c.alignment = Alignment(vertical="center")
for r in ROWS:
    ws.append(list(r))
for row in ws.iter_rows(min_row=5, max_row=4 + len(ROWS)):
    row[4].number_format = "DD.MM.YYYY"
    row[5].number_format = "0%"
    row[6].number_format = "DD.MM.YYYY"
widths = [24, 18, 12, 38, 16, 13, 16]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:G{4 + len(ROWS)}"
wb.save(Path(__file__).parent / "team" / "team_list.xlsx")
print("ok")
