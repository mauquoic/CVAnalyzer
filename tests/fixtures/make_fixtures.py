"""Generate synthetic (fictional) PowerPoint CVs and an Excel team list for testing the tools.

Needs python-pptx and openpyxl (dev only): pip install python-pptx openpyxl
Usage: python tests/fixtures/make_fixtures.py
"""
import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches, Pt

HERE = Path(__file__).parent

CVS = {
    # file name: (heading, slides)
    "Anna_Mueller_CV_2026-02.pptx": ("Anna Müller – Senior Fullstack Developer", [
        ["Profile", "12 years of experience in software development, of which 9 years as fullstack developer.",
         "Languages: German (native), English (fluent, C1)"],
        ["Skills", "Java, Spring Boot, Python", "TypeScript, Angular 15", "Apache Spark, PySpark, Databricks",
         "REST APIs, OpenAPI", "Docker, Kubernetes, Azure", "GitLab, merge requests, code reviews",
         "GitHub Copilot for test generation"],
        ["Projects", "2021–2025 Insurance client: built data pipelines (ETL) with PySpark on Databricks and an Angular front end.",
         "2016–2021 Bank: Java microservices with REST APIs, lead developer for a team of 5."],
    ]),
    "Anna_Mueller_CV_2023-05.pptx": ("Anna Müller – Fullstack Developer", [
        ["Skills", "Java, Angular"],
    ]),
    "CV Jonas Peter Becker.pptx": ("Jonas Becker – Fullstack Developer", [
        ["Profile", "6 years of experience as a developer.", "Languages: English (fluent), German (basic, A2)"],
        ["Skills", "Python, Django, FastAPI", "TypeScript, Angular", "REST APIs", "Git, pull requests, code reviews",
         "Docker, AWS"],
        ["Projects", "2020–2026 Retail client: Python backend with REST APIs and Angular front end, ETL jobs with Airflow."],
    ]),
    "Weiss_Joerg.pptx": ("Jörg Weiß – Data Engineer", [
        ["Profile", "8 years in data engineering.", "Languages: German (native), English (fluent)"],
        ["Skills", "Python, Apache Spark, Spark SQL", "Airflow, dbt, ETL/ELT", "Kubernetes, GCP"],
        ["Projects", "2018–2026 Telecom: data pipeline development with Spark SQL and Airflow."],
    ]),
    "Jana_Schmidt.pptx": ("Jana Schmidt – Frontend Developer", [
        ["Profile", "5 years of frontend development.", "Languages: German (native), English (good)"],
        ["Skills", "TypeScript, Angular, React", "REST API integration", "Git"],
        ["Projects", "2021–2026 Public sector: Angular applications consuming REST APIs."],
    ]),
    "Unknown_Person_CV.pptx": ("Max Mustermann – Developer", [
        ["Skills", "C#, .NET, Azure", "Languages: German, English"],
    ]),
}

TEAM = [
    ["Name", "Level", "Verfügbarkeit", "Verfügbar ab"],
    ["Anna Maria Müller", "Senior", "0%", "2026-11-01"],
    ["Jonas Becker", "Consultant", "100%", ""],
    ["Jörg Weiss", "Senior Consultant", "50%", ""],
    ["Jan Schmidt", "Consultant", "100%", ""],
    ["Lena Vogel", "Consultant", "100%", ""],
]


def make_pptx(path: Path, heading: str, slides: list[list[str]], saved: dt.datetime) -> None:
    prs = Presentation()
    prs.core_properties.modified = saved
    blank = prs.slide_layouts[6]
    for i, lines in enumerate(slides):
        slide = prs.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(6)).text_frame
        box.text = heading if i == 0 else lines[0]
        box.paragraphs[0].runs[0].font.size = Pt(24)
        for line in (lines if i == 0 else lines[1:]):
            box.add_paragraph().text = line
    prs.save(path)


if __name__ == "__main__":
    for name, (heading, slides) in CVS.items():
        # Jörg's CV is deliberately old, to exercise the outdated flag.
        saved = dt.datetime(2024, 3, 15) if name.startswith("Weiss") else dt.datetime(2026, 6, 1)
        make_pptx(HERE / "cvs" / name, heading, slides, saved)
    wb = Workbook()
    ws = wb.active
    ws.append(["Team Data & Engineering (synthetic test data)"])
    for row in TEAM:
        ws.append(row)
    wb.save(HERE / "team" / "team.xlsx")
    print("Fixtures written to", HERE)
